import os, calendar, io, unicodedata, re
from datetime import datetime, timedelta, date
from flask import Flask, render_template, redirect, url_for, request, session, flash, send_file
from flask_sqlalchemy import SQLAlchemy
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# Cấu hình đường dẫn cơ sở dữ liệu SQLite
basedir = os.path.abspath(os.path.dirname(__file__))
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + os.path.join(basedir, 'quanly.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.secret_key = 'super_secret_key_change_this'

db = SQLAlchemy(app)

# ==============================================================================
# 1. CÁC HÀM TIỆN ÍCH (HELPERS)
# ==============================================================================
AVATAR_COLORS = [
    '#10b981', '#f59e0b', '#3b82f6', '#8b5cf6', '#ec4899', 
    '#06b6d4', '#14b8a6', '#f97316', '#6366f1', '#84cc16',
    '#d97706', '#0284c7', '#7c3aed', '#db2777'
]

def get_avatar_color(text):
    if not text:
        return '#3b82f6'
    return AVATAR_COLORS[sum(ord(c) for c in str(text)) % len(AVATAR_COLORS)]

def get_initials(name):
    if not name:
        return 'NV'
    parts = name.strip().split()
    if len(parts) >= 2:
        return (parts[0][0] + parts[-1][0]).upper()
    return name[:2].upper()

def remove_accents(text):
    if not text:
        return ''
    s1 = unicodedata.normalize('NFD', text)
    s2 = ''.join(c for c in s1 if unicodedata.category(c) != 'Mn')
    s2 = s2.replace('đ', 'd').replace('Đ', 'D')
    return s2

# Định dạng bảng Excel chuẩn mực, tinh gọn, không màu mè
EXCEL_FONT_TITLE = Font(name='Segoe UI', size=12, bold=True)
EXCEL_FONT_SUB = Font(name='Segoe UI', size=9, italic=True, color='4B5563')
EXCEL_FONT_HEADER = Font(name='Segoe UI', size=10, bold=True)
EXCEL_FONT_DATA = Font(name='Segoe UI', size=9.5)
EXCEL_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)
EXCEL_ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
EXCEL_ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)

def write_simple_table(ws, title, subtitle, headers, data_rows):
    """Ghi bảng Excel thông tin sạch sẽ, không màu mè, tự căn chỉnh độ rộng cột."""
    ws.cell(row=1, column=1, value=title).font = EXCEL_FONT_TITLE
    if subtitle:
        ws.cell(row=2, column=1, value=subtitle).font = EXCEL_FONT_SUB
    
    header_row = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = EXCEL_FONT_HEADER
        cell.alignment = EXCEL_ALIGN_CENTER
        cell.border = EXCEL_BORDER

    for r_idx, row_items in enumerate(data_rows, start=header_row + 1):
        for col_idx, item in enumerate(row_items, start=1):
            val = item[0] if isinstance(item, tuple) else item
            align = item[1] if isinstance(item, tuple) else EXCEL_ALIGN_LEFT
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = EXCEL_FONT_DATA
            cell.alignment = align
            cell.border = EXCEL_BORDER

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max((len(str(c.value or '').split('\n')[0]) for c in col if c.row >= header_row), default=0)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

# ==============================================================================
# 2. ĐỊNH NGHĨA CƠ SỞ DỮ LIỆU (DATABASE MODELS)
# ==============================================================================

# 1. Bảng Employee (Nhân viên)
class Employee(db.Model):
    __tablename__ = 'employee'
    
    id = db.Column(db.Integer, primary_key=True)
    employee_code = db.Column(db.String(50), unique=True, nullable=False)
    full_name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(100))
    phone = db.Column(db.String(20))
    emergency_phone = db.Column(db.String(50))
    position = db.Column(db.String(100))
    department = db.Column(db.String(50), nullable=False, default='TDH')  # 'TDH', 'A11A', 'LENS'
    role = db.Column(db.String(20), nullable=False, default='user')        # 'admin' hoặc 'user'
    password_hash = db.Column(db.String(200), nullable=False)
    start_date = db.Column(db.Date)
    status = db.Column(db.String(50), default='active')                    # 'active' hoặc 'inactive'

# 2. Bảng Task (Nhiệm vụ / Dự án)
class Task(db.Model):
    __tablename__ = 'task'
    
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    employee_code = db.Column(db.String(50), db.ForeignKey('employee.employee_code'), nullable=False)
    deadline = db.Column(db.Date)
    priority = db.Column(db.String(50), default='Bình thường')
    description = db.Column(db.Text)
    status = db.Column(db.String(50), default='Đang thực hiện')
    progress = db.Column(db.Integer, default=100)
    rating = db.Column(db.String(50), default='Tốt')

    # Liên kết với bảng Employee
    assignee = db.relationship('Employee', backref=db.backref('tasks', lazy=True))

# 3. Bảng LeaveRequest (Đăng ký lịch nghỉ phép)
class LeaveRequest(db.Model):
    __tablename__ = 'leave_request'
    
    id = db.Column(db.Integer, primary_key=True)
    employee_code = db.Column(db.String(50), db.ForeignKey('employee.employee_code'), nullable=False)
    leave_date = db.Column(db.Date, nullable=False)
    reason = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.now)

    # Liên kết với bảng Employee
    employee = db.relationship('Employee', backref=db.backref('leave_requests', lazy=True))

# 4. Bảng Report (Báo cáo công việc hàng ngày)
class Report(db.Model):
    __tablename__ = 'report'
    
    id = db.Column(db.Integer, primary_key=True)
    employee_code = db.Column(db.String(50), db.ForeignKey('employee.employee_code'), nullable=False)
    report_date = db.Column(db.Date, nullable=False)
    work_types = db.Column(db.String(255))
    description = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.now)

    # Liên kết với bảng Employee
    employee = db.relationship('Employee', backref=db.backref('reports', lazy=True))

# 5. Bảng DepartmentNotice (Cài đặt thông báo phòng ban A11A, LENS)
class DepartmentNotice(db.Model):
    __tablename__ = 'department_notice'
    
    id = db.Column(db.Integer, primary_key=True)
    department = db.Column(db.String(50), unique=True, nullable=False)  # 'A11A', 'LENS'
    content = db.Column(db.Text, default='')
    is_active = db.Column(db.Boolean, default=False)
    updated_at = db.Column(db.DateTime, default=datetime.now, onupdate=datetime.now)

# 6. Bảng ProjectGroup (Quản lý Dự án Chung - Tạo nhóm và gom các dự án nhỏ)
class ProjectGroup(db.Model):
    __tablename__ = 'project_group'
    
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(255), nullable=False)
    department = db.Column(db.String(50), nullable=False, default='TDH')
    creator_name = db.Column(db.String(100), default='Quản Trị Viên')
    members_str = db.Column(db.String(255), default='')
    task_count = db.Column(db.Integer, default=1)
    created_at = db.Column(db.DateTime, default=datetime.now)

# ==============================================================================
# 3. ROUTE XÁC THỰC (AUTHENTICATION)
# ==============================================================================

@app.route('/', methods=['GET', 'POST'])
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        emp_code = request.form.get('employee_code', '').strip()
        password = request.form.get('password', '').strip()
        
        employee = Employee.query.filter_by(employee_code=emp_code).first()
        
        if employee and employee.password_hash == password:
            session['employee_id'] = employee.id
            session['employee_code'] = employee.employee_code
            session['role'] = employee.role
            session['full_name'] = employee.full_name
            session['department'] = employee.department or 'TDH'
            session['notice_seen'] = False
            
            if employee.role == 'admin':
                return redirect(url_for('admin_dashboard'))
            else:
                return redirect(url_for('user_tasks'))
        else:
            flash('Mã nhân viên hoặc mật khẩu không chính xác!', 'danger')
            
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/forgot-password', methods=['POST'])
def forgot_password():
    email = request.form.get('email')
    flash('Đã gửi yêu cầu cho quản trị viên. Hãy kiểm tra mail của bạn', 'success')
    return redirect(url_for('login'))

@app.route('/dashboard')
def user_dashboard():
    if not session.get('employee_code'):
        return redirect(url_for('login'))
    if session.get('role') == 'admin':
        return redirect(url_for('admin_dashboard'))
    return redirect(url_for('user_tasks'))

# ==============================================================================
# 4. ROUTE DÀNH CHO USER: CÔNG VIỆC CỦA TÔI
# ==============================================================================

@app.route('/user/tasks')
def user_tasks():
    emp_code = session.get('employee_code')
    if not emp_code:
        flash('Vui lòng đăng nhập trước khi tiếp tục.', 'warning')
        return redirect(url_for('login'))
        
    user = Employee.query.filter_by(employee_code=emp_code).first()
    if not user:
        session.clear()
        return redirect(url_for('login'))

    is_admin = (session.get('role') == 'admin')

    # Lấy danh sách nhiệm vụ của riêng user đang đăng nhập
    tasks = Task.query.filter_by(employee_code=emp_code).order_by(Task.id.desc()).all()
    
    # Tính toán các chỉ số thống kê cá nhân động từ CSDL
    total_tasks = len(tasks)
    in_progress_tasks = sum(1 for t in tasks if t.status == 'Đang thực hiện')
    completed_tasks = sum(1 for t in tasks if t.status == 'Đã hoàn thành')
    stuck_tasks = sum(1 for t in tasks if t.status in ['Vướng mắc', 'Chưa hoàn thành'])

    # Lấy danh sách báo cáo trong ngày hôm nay của riêng user (ngày nào hiển thị ngày đó, báo cáo cũ không hiển thị)
    today_date = date.today()

    reports = Report.query.filter(
        Report.employee_code == emp_code,
        Report.report_date == today_date
    ).order_by(Report.id.desc()).all()

    for r in reports:
        # Báo cáo trong ngày hôm nay có thể chỉnh sửa/xóa bởi chính user
        r.can_edit = True

    initials = get_initials(user.full_name)
    today_str = datetime.now().strftime('%Y-%m-%d')

    # Lấy thông báo nếu thuộc phòng A11A hoặc LENS và đang bật (chỉ xuất hiện 1 lần đầu tiên khi đăng nhập)
    dept_notice = None
    show_notice = False
    if user.department in ['A11A', 'LENS']:
        dept_notice = DepartmentNotice.query.filter_by(department=user.department, is_active=True).first()
        if dept_notice and dept_notice.is_active and dept_notice.content:
            if not session.get('notice_seen'):
                show_notice = True
                session['notice_seen'] = True

    return render_template(
        'user_tasks.html',
        user=user,
        tasks=tasks,
        total_tasks=total_tasks,
        in_progress_tasks=in_progress_tasks,
        completed_tasks=completed_tasks,
        stuck_tasks=stuck_tasks,
        reports=reports,
        dept_notice=dept_notice,
        show_notice=show_notice,
        initials=initials,
        today_str=today_str,
        is_admin=is_admin,
        active_page='user_tasks'
    )

# Tạo dự án / công việc mới (User Modal)
@app.route('/user/tasks/create', methods=['POST'])
def user_create_task():
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    title = request.form.get('title', '').strip()
    deadline_str = request.form.get('deadline')
    priority = request.form.get('priority', 'Bình thường')
    description = request.form.get('description', '').strip()

    if not title:
        flash('Vui lòng nhập tên công việc / dự án!', 'warning')
        return redirect(request.referrer or url_for('user_tasks'))

    deadline = None
    if deadline_str:
        try:
            deadline = datetime.strptime(deadline_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    new_task = Task(
        title=title,
        employee_code=emp_code,
        deadline=deadline,
        priority=priority,
        description=description,
        status='Đang thực hiện',
        progress=0,
        rating='—'
    )
    db.session.add(new_task)
    db.session.commit()
    flash(f'Đã tạo công việc / dự án "{title}" thành công!', 'success')
    return redirect(request.referrer or url_for('user_tasks'))

# Báo cáo công việc hàng ngày (User Modal)
@app.route('/user/report/create', methods=['POST'])
def user_create_report():
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    report_date_str = request.form.get('report_date')
    work_types_list = request.form.getlist('work_types')
    work_types_str = ", ".join(work_types_list) if work_types_list else "Khác"
    description = request.form.get('description', '').strip()

    if not description:
        flash('Vui lòng nhập mô tả chi tiết công việc đã thực hiện!', 'warning')
        return redirect(request.referrer or url_for('user_tasks'))

    report_date = datetime.now().date()
    if report_date_str:
        try:
            report_date = datetime.strptime(report_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    new_report = Report(
        employee_code=emp_code,
        report_date=report_date,
        work_types=work_types_str,
        description=description,
        created_at=datetime.now()
    )
    db.session.add(new_report)
    db.session.commit()
    flash('Gửi báo cáo công việc hàng ngày thành công!', 'success')
    return redirect(request.referrer or url_for('user_tasks'))

# Cập nhật nhanh trạng thái công việc
@app.route('/user/tasks/update-status/<int:id>', methods=['POST'])
def user_update_task_status(id):
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    task = Task.query.get_or_404(id)
    if task.employee_code == emp_code or session.get('role') == 'admin':
        new_status = request.form.get('status')
        if new_status in ['Đang thực hiện', 'Đã hoàn thành', 'Vướng mắc']:
            task.status = new_status
            if new_status == 'Đã hoàn thành':
                task.progress = 100
            elif new_status == 'Vướng mắc':
                task.progress = 45
            db.session.commit()
            flash(f'Đã cập nhật trạng thái công việc sang "{new_status}"!', 'success')
    else:
        flash('Bạn không có quyền cập nhật công việc này!', 'danger')
        
    return redirect(request.referrer or url_for('user_tasks'))

# Xóa công việc cá nhân
@app.route('/user/tasks/delete/<int:id>', methods=['POST'])
def user_delete_task(id):
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    task = Task.query.get_or_404(id)
    if task.employee_code == emp_code or session.get('role') == 'admin':
        db.session.delete(task)
        db.session.commit()
        flash('Đã xóa công việc thành công!', 'info')
    else:
        flash('Bạn không có quyền xóa công việc này!', 'danger')
        
    return redirect(request.referrer or url_for('user_tasks'))

# Chỉnh sửa chi tiết công việc cá nhân
@app.route('/user/tasks/edit/<int:id>', methods=['POST'])
def user_edit_task(id):
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    task = Task.query.get_or_404(id)
    if task.employee_code == emp_code or session.get('role') == 'admin':
        title = request.form.get('title', '').strip()
        deadline_str = request.form.get('deadline')
        priority = request.form.get('priority', 'Bình thường')
        status = request.form.get('status', 'Đang thực hiện')
        description = request.form.get('description', '').strip()

        if title:
            task.title = title
        if deadline_str:
            try:
                task.deadline = datetime.strptime(deadline_str, '%Y-%m-%d').date()
            except ValueError:
                pass
        elif 'deadline' in request.form:
            task.deadline = None

        task.priority = priority
        task.status = status
        if status == 'Đã hoàn thành':
            task.progress = 100
        elif status == 'Vướng mắc':
            task.progress = 45
        
        task.description = description
        db.session.commit()
        flash(f'Đã cập nhật công việc "{task.title}" thành công!', 'success')
    else:
        flash('Bạn không có quyền chỉnh sửa công việc này!', 'danger')
        
    return redirect(request.referrer or url_for('user_tasks'))

# Chỉnh sửa báo cáo công việc cá nhân
@app.route('/user/report/edit/<int:id>', methods=['POST'])
def user_edit_report(id):
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    report = Report.query.get_or_404(id)
    if report.employee_code != emp_code and session.get('role') != 'admin':
        flash('Bạn không có quyền chỉnh sửa báo cáo này!', 'danger')
        return redirect(request.referrer or url_for('user_tasks'))
        
    # Chỉ cho phép chỉnh sửa báo cáo trong ngày (báo cáo cũ không thể sửa)
    today_date = date.today()
    is_today = (report.report_date == today_date or (report.created_at and report.created_at.date() == today_date))
    if not is_today and session.get('role') != 'admin':
        flash('Không thể chỉnh sửa báo cáo cũ! Bạn chỉ có thể sửa báo cáo trong ngày.', 'danger')
        return redirect(request.referrer or url_for('user_tasks'))

    report_date_str = request.form.get('report_date')
    work_types_list = request.form.getlist('work_types')
    work_types_str = ", ".join(work_types_list) if work_types_list else "Khác"
    description = request.form.get('description', '').strip()

    if not description:
        flash('Vui lòng nhập mô tả chi tiết công việc đã thực hiện!', 'warning')
        return redirect(request.referrer or url_for('user_tasks'))

    if report_date_str:
        try:
            parsed_date = datetime.strptime(report_date_str, '%Y-%m-%d').date()
            if parsed_date == today_date or session.get('role') == 'admin':
                report.report_date = parsed_date
        except ValueError:
            pass

    report.work_types = work_types_str
    report.description = description
    db.session.commit()
    flash('Cập nhật báo cáo công việc thành công!', 'success')
    return redirect(request.referrer or url_for('user_tasks'))

# Xóa báo cáo công việc cá nhân
@app.route('/user/report/delete/<int:id>', methods=['POST'])
def user_delete_report(id):
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    report = Report.query.get_or_404(id)
    if report.employee_code != emp_code and session.get('role') != 'admin':
        flash('Bạn không có quyền xóa báo cáo này!', 'danger')
        return redirect(request.referrer or url_for('user_tasks'))

    # Chỉ cho phép xóa báo cáo trong ngày
    today_date = date.today()
    is_today = (report.report_date == today_date or (report.created_at and report.created_at.date() == today_date))
    if not is_today and session.get('role') != 'admin':
        flash('Không thể xóa báo cáo cũ! Bạn chỉ có thể xóa báo cáo trong ngày.', 'danger')
        return redirect(request.referrer or url_for('user_tasks'))

    db.session.delete(report)
    db.session.commit()
    flash('Đã xóa báo cáo công việc thành công!', 'info')
    return redirect(request.referrer or url_for('user_tasks'))
        
    return redirect(request.referrer or url_for('user_tasks'))

# ==============================================================================
# 5. ROUTE QUẢN LÝ DỰ ÁN (DÙNG CHUNG ADMIN & USER VỚI BỘ LỌC PHÒNG BAN)
# ==============================================================================

@app.route('/admin')
@app.route('/projects')
def admin_dashboard():
    if not session.get('employee_code'):
        return redirect(url_for('login'))
    
    current_user = Employee.query.filter_by(employee_code=session.get('employee_code')).first()
    if not current_user:
        session.clear()
        return redirect(url_for('login'))
        
    is_admin = (session.get('role') == 'admin')
    if is_admin:
        selected_dept = request.args.get('dept', 'TDH')
        if selected_dept not in ['TDH', 'A11A', 'LENS']:
            selected_dept = 'TDH'
    else:
        # User thường: Tự động khóa và hiển thị theo đúng department của họ
        user_dept = current_user.department or 'TDH'
        requested_dept = request.args.get('dept')
        if requested_dept and requested_dept != user_dept:
            flash(f'Bạn chỉ có quyền xem dữ liệu phòng ban của mình ({user_dept})!', 'info')
            return redirect(url_for('admin_dashboard', dept=user_dept))
        selected_dept = user_dept

    # Lọc nhân viên và công việc chỉ theo 3 phòng ban cụ thể (TDH, A11A, LENS) - chỉ lấy nhân viên đang làm việc
    employees = Employee.query.filter(
        Employee.department == selected_dept,
        (Employee.status != 'inactive') | (Employee.status == None)
    ).all()
    tasks = Task.query.join(Employee, Task.employee_code == Employee.employee_code)\
                      .filter(Employee.department == selected_dept).all()
        
    all_employees = employees
    
    # Gán initials, avatar_color và thống kê cho từng nhân viên
    for emp in employees:
        emp.initials = get_initials(emp.full_name)
        emp.avatar_color = get_avatar_color(emp.employee_code + emp.full_name)
        emp_tasks = [t for t in tasks if t.employee_code == emp.employee_code]
        emp.total_tasks = len(emp_tasks)
        emp.completed_tasks = sum(1 for t in emp_tasks if t.status == 'Đã hoàn thành')
        ratings = [getattr(t, 'rating', None) for t in emp_tasks if getattr(t, 'rating', None) and t.rating != '-']
        emp.avg_rating = ratings[0] if ratings else ('Tốt' if emp.completed_tasks > 0 else '—')
        emp.is_leader = ('Leader' in (emp.position or ''))

    # Thống kê nhanh hoàn toàn động từ DB
    total_tasks = len(tasks)
    in_progress_tasks = sum(1 for t in tasks if t.status == 'Đang thực hiện')
    completed_tasks = sum(1 for t in tasks if t.status == 'Đã hoàn thành')
    stuck_tasks = sum(1 for t in tasks if t.status in ['Vướng mắc', 'Chưa hoàn thành'])

    completed_pct = round((completed_tasks / total_tasks * 100), 1) if total_tasks else 0
    stuck_pct = round((stuck_tasks / total_tasks * 100), 1) if total_tasks else 0
    in_progress_pct = round((in_progress_tasks / total_tasks * 100), 1) if total_tasks else 0

    today = date.today()
    for t in tasks:
        emp = next((e for e in all_employees if e.employee_code == t.employee_code), None)
        t.emp_name = emp.full_name if emp else (t.employee_code or 'Nhân viên')
        t.emp_initials = emp.initials if emp else 'NV'
        t.emp_avatar_color = getattr(emp, 'avatar_color', '#2563eb')
        t.deadline_str = t.deadline.strftime('%d/%m/%Y') if t.deadline else '—'
        t.is_overdue = (t.deadline < today and t.status != 'Đã hoàn thành') if t.deadline else False

        if t.progress is not None:
            t.pct = t.progress
        elif t.status == 'Đã hoàn thành':
            t.pct = 100
        elif t.status == 'Vướng mắc':
            t.pct = 45
        else:
            t.pct = 0

        if t.pct == 100 or t.status == 'Đã hoàn thành':
            t.pct_color = '#10b981'
        elif t.status == 'Vướng mắc':
            t.pct_color = '#ef4444'
        else:
            t.pct_color = '#2563eb'

    # Dữ liệu cho biểu đồ tiến độ dự án từ DB
    chart_bar_labels = []
    chart_bar_data = []
    chart_bar_colors = []
    chart_project_details = []
    
    for t in tasks[:20]:
        chart_bar_labels.append(t.title[:22] + '...' if len(t.title) > 22 else t.title)
        pct = t.pct
        chart_bar_data.append(pct)
        if pct == 100 or t.status == 'Đã hoàn thành':
            color = '#10b981'
            status_text = 'HOÀN THÀNH'
        elif t.status == 'Vướng mắc':
            color = '#ef4444'
            status_text = 'VƯỚNG MẮC'
        else:
            color = '#f59e0b'
            status_text = 'ĐANG CHẠY'
            
        chart_bar_colors.append(color)
        chart_project_details.append({
            'title': t.title,
            'emp_name': t.emp_name,
            'status_text': status_text,
            'status_color': color,
            'pct': pct
        })

    # Dữ liệu biểu đồ 7 ngày qua cho tab 'Tiến độ'
    days_7 = [today - timedelta(days=6 - i) for i in range(7)]
    timeline_labels = [d.strftime('%d/%m') for d in days_7]
    timeline_totals = []
    timeline_details = []

    for d in days_7:
        day_tasks = [t for t in tasks if t.deadline == d]
        day_total = 0.0
        emp_breakdown = {}
        
        for t in day_tasks:
            pct_val = float(t.pct or 0)
            if pct_val > 0:
                day_total += pct_val
                if t.emp_name not in emp_breakdown:
                    emp_breakdown[t.emp_name] = {'total': 0.0, 'tasks': []}
                emp_breakdown[t.emp_name]['total'] += pct_val
                emp_breakdown[t.emp_name]['tasks'].append({
                    'title': t.title,
                    'pct': int(pct_val)
                })

        timeline_totals.append(round(day_total, 1))
        timeline_details.append({
            'date_str': d.strftime('%d/%m'),
            'total': round(day_total, 1),
            'employees': emp_breakdown
        })

    # Legend danh sách thành viên nhóm dự án
    seen_names = set()
    team_members_legend = []
    palette = [
        '#ec4899', '#f59e0b', '#8b5cf6', '#f43f5e', '#f97316',
        '#a855f7', '#06b6d4', '#eab308', '#3b82f6', '#10b981',
        '#14b8a6', '#6366f1', '#22c55e', '#d946ef', '#2563eb', '#84cc16'
    ]
    p_idx = 0
    for emp in employees:
        if emp.full_name not in seen_names:
            seen_names.add(emp.full_name)
            color = palette[p_idx % len(palette)]
            p_idx += 1
            team_members_legend.append({
                'name': emp.full_name,
                'color': color
            })

    all_tasks_json = [
        {
            'id': t.id,
            'title': t.title,
            'employee_code': t.employee_code,
            'priority': t.priority,
            'deadline': t.deadline.strftime('%d/%m/%Y') if t.deadline else None,
            'status': t.status
        } for t in tasks if is_admin or t.employee_code == current_user.employee_code
    ]

    admin_initials = get_initials(current_user.full_name)

    # Thông báo phòng ban cho A11A và LENS (Admin thấy 1 lần duy nhất khi chuyển sang nhóm)
    dept_notice = None
    show_admin_notice = False
    if selected_dept in ['A11A', 'LENS']:
        dept_notice = DepartmentNotice.query.filter_by(department=selected_dept).first()
        if is_admin and dept_notice and dept_notice.is_active and dept_notice.content:
            notice_key = f'admin_notice_seen_{selected_dept}'
            if not session.get(notice_key):
                show_admin_notice = True
                session[notice_key] = True

    # Danh sách nhóm dự án chung (Dành cho modal bánh răng Quản lý Dự án Chung)
    project_groups = ProjectGroup.query.filter_by(department=selected_dept).order_by(ProjectGroup.id.asc()).all()

    return render_template(
        'admin_dashboard.html',
        employees=employees,
        all_employees=all_employees,
        tasks=tasks,
        selected_dept=selected_dept,
        total_tasks=total_tasks,
        in_progress_tasks=in_progress_tasks,
        completed_tasks=completed_tasks,
        stuck_tasks=stuck_tasks,
        completed_pct=completed_pct,
        stuck_pct=stuck_pct,
        in_progress_pct=in_progress_pct,
        chart_bar_labels=chart_bar_labels,
        chart_bar_data=chart_bar_data,
        chart_bar_colors=chart_bar_colors,
        chart_project_details=chart_project_details,
        timeline_labels=timeline_labels,
        timeline_totals=timeline_totals,
        timeline_details=timeline_details,
        team_members_legend=team_members_legend,
        project_groups=project_groups,
        all_tasks_json=all_tasks_json,
        dept_notice=dept_notice,
        show_admin_notice=show_admin_notice,
        admin_initials=admin_initials,
        current_user=current_user,
        is_admin=is_admin,
        active_page='projects'
    )

# ==============================================================================
# QUẢN LÝ DỰ ÁN CHUNG (MODAL BÁNH RĂNG CẠNH NĂM 2026)
# ==============================================================================

@app.route('/admin/project-group/save', methods=['POST'])
def save_project_group():
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    current_user = Employee.query.filter_by(employee_code=session.get('employee_code')).first()
    if not current_user:
        return redirect(url_for('login'))
        
    group_id = request.form.get('group_id')
    name = request.form.get('name', '').strip()
    members_str = request.form.get('members_str', '').strip()
    task_count = request.form.get('task_count', 1, type=int)
    dept = request.form.get('dept', current_user.department or 'TDH')

    if not name:
        flash('Vui lòng nhập tên nhóm dự án!', 'warning')
        return redirect(url_for('admin_dashboard', dept=dept))

    if group_id:
        group = ProjectGroup.query.get(group_id)
        if group:
            group.name = name
            if members_str:
                group.members_str = members_str
            if task_count:
                group.task_count = task_count
            db.session.commit()
            flash(f'Đã cập nhật nhóm dự án "{name}" thành công!', 'success')
    else:
        new_group = ProjectGroup(
            name=name,
            department=dept,
            creator_name=current_user.full_name,
            members_str=members_str or current_user.full_name,
            task_count=task_count or 1
        )
        db.session.add(new_group)
        db.session.commit()
        flash(f'Đã tạo nhóm dự án "{name}" thành công!', 'success')

    return redirect(url_for('admin_dashboard', dept=dept))

@app.route('/admin/project-group/delete/<int:id>', methods=['POST'])
def delete_project_group(id):
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    group = ProjectGroup.query.get_or_404(id)
    dept = group.department or 'TDH'
    group_name = group.name
    db.session.delete(group)
    db.session.commit()
    flash(f'Đã xóa nhóm dự án "{group_name}"!', 'info')
    return redirect(url_for('admin_dashboard', dept=dept))

# Cài đặt thông báo phòng ban (Chỉ dành cho Admin, áp dụng cho A11A và LENS)
@app.route('/admin/notice/save', methods=['POST'])
def save_department_notice():
    if not session.get('employee_code') or session.get('role') != 'admin':
        flash('Chỉ Quản trị viên mới có quyền cài đặt thông báo!', 'danger')
        return redirect(url_for('admin_dashboard'))
        
    dept = request.form.get('dept', 'A11A')
    if dept not in ['A11A', 'LENS']:
        dept = 'A11A'
        
    content = request.form.get('content', '').strip()
    is_active = (request.form.get('is_active') == 'on' or request.form.get('is_active') == 'true')
    
    notice = DepartmentNotice.query.filter_by(department=dept).first()
    if not notice:
        notice = DepartmentNotice(department=dept, content=content, is_active=is_active)
        db.session.add(notice)
    else:
        notice.content = content
        notice.is_active = is_active
        
    db.session.commit()
    flash(f'Đã lưu cài đặt thông báo cho nhóm {dept} thành công!', 'success')
    return redirect(url_for('admin_dashboard', dept=dept))

# Giao việc mới (Từ Modal trong Quản lý dự án)
@app.route('/admin/assign-task', methods=['GET', 'POST'])
def assign_task():
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        employee_code = request.form.get('employee_code')
        deadline_str = request.form.get('deadline')
        priority = request.form.get('priority', 'Bình thường')
        description = request.form.get('description', '')
        dept = request.form.get('dept', '')
        progress = request.form.get('progress', 0, type=int) if request.form.get('progress') else 0
        rating = request.form.get('rating', '—') or '—'
        status = request.form.get('status', 'Đang thực hiện')
        
        deadline = None
        if deadline_str:
            try:
                deadline = datetime.strptime(deadline_str, '%Y-%m-%d').date()
            except ValueError:
                pass
            
        new_task = Task(
            title=title,
            employee_code=employee_code,
            deadline=deadline,
            priority=priority,
            description=description,
            status=status,
            progress=progress,
            rating=rating
        )
        db.session.add(new_task)
        db.session.commit()
        
        flash('Giao việc thành công!', 'success')
        return redirect(url_for('admin_dashboard', dept=dept) if dept else url_for('admin_dashboard'))
        
    return redirect(url_for('admin_dashboard'))

# ==============================================================================
# 6. ROUTE QUẢN LÝ NHÂN SỰ (ADMIN TOÀN QUYỀN - USER CHỈ TỰ SỬA CÁ NHÂN)
# ==============================================================================

@app.route('/admin/employees')
def admin_employees():
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    current_user = Employee.query.filter_by(employee_code=session.get('employee_code')).first()
    if not current_user:
        session.clear()
        return redirect(url_for('login'))

    is_admin = (session.get('role') == 'admin')
    
    if is_admin:
        selected_dept = request.args.get('dept', 'TDH')
        if selected_dept not in ['TDH', 'A11A', 'LENS']:
            selected_dept = 'TDH'
    else:
        user_dept = current_user.department or 'TDH'
        requested_dept = request.args.get('dept')
        if requested_dept and requested_dept != user_dept:
            flash(f'Bạn chỉ có quyền xem danh sách nhân sự phòng ban của mình ({user_dept})!', 'info')
            return redirect(url_for('admin_employees', dept=user_dept))
        selected_dept = user_dept

    current_tab = request.args.get('tab', 'active')
    if current_tab not in ['active', 'inactive']:
        current_tab = 'active'

    # Truy vấn động nhân viên từ DB theo phòng ban (TDH, A11A, LENS)
    all_dept_employees = Employee.query.filter_by(department=selected_dept).all()

    active_employees = []
    inactive_employees = []

    # Tính toán số liệu thống kê dự án của từng nhân sự từ CSDL
    for emp in all_dept_employees:
        emp_tasks = Task.query.filter_by(employee_code=emp.employee_code).all()
        emp.total_tasks = len(emp_tasks)
        emp.completed_tasks = sum(1 for t in emp_tasks if t.status == 'Đã hoàn thành')
        emp.on_time_pct = round((emp.completed_tasks / emp.total_tasks * 100)) if emp.total_tasks > 0 else 0
        emp.initials = get_initials(emp.full_name)
        emp.avatar_color = get_avatar_color(emp.employee_code + emp.full_name)
        
        pos = (emp.position or '').strip()
        emp.position_short = (pos[:11] + '..') if len(pos) > 13 else pos
        emp.start_date_formatted = emp.start_date.strftime('%d/%m/%Y') if emp.start_date else 'Chưa cập nhật'

        if emp.status == 'inactive':
            inactive_employees.append(emp)
        else:
            active_employees.append(emp)

    admin_initials = get_initials(current_user.full_name)
    return render_template(
        'admin_employees.html',
        active_employees=active_employees,
        inactive_employees=inactive_employees,
        employees=active_employees,  # Tương thích ngược nếu còn biến cũ
        current_tab=current_tab,
        selected_dept=selected_dept,
        admin_initials=admin_initials,
        current_user=current_user,
        is_admin=is_admin,
        active_page='employees'
    )

# Thêm nhân sự mới (Chỉ Admin)
@app.route('/admin/add', methods=['POST'])
def add_employee():
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    if session.get('role') != 'admin':
        flash('Chỉ Quản trị viên (Admin) mới có quyền thêm nhân sự mới!', 'danger')
        return redirect(url_for('admin_employees'))
        
    code = request.form.get('employee_code', '').strip()
    raw_password = request.form.get('password', '123456').strip()
    role = request.form.get('role', 'user')
    name = request.form.get('full_name', '').strip()
    email = request.form.get('email', '').strip()
    phone = request.form.get('phone', '').strip()
    emergency_phone = request.form.get('emergency_phone', '').strip()
    position = request.form.get('position', '').strip()
    department = request.form.get('department', 'TDH')
    start_date_str = request.form.get('start_date')
    start_date = None
    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    
    if Employee.query.filter_by(employee_code=code).first():
        flash(f'Mã nhân viên {code} đã tồn tại trong hệ thống!', 'danger')
        return redirect(url_for('admin_employees', dept=department))
        
    new_emp = Employee(
        employee_code=code,
        password_hash=raw_password,
        role=role,
        full_name=name,
        email=email,
        phone=phone,
        emergency_phone=emergency_phone,
        position=position,
        department=department,
        start_date=start_date,
        status='active'
    )
    db.session.add(new_emp)
    db.session.commit()
    flash(f'Thêm nhân sự {name} ({code}) thành công!', 'success')
    return redirect(url_for('admin_employees', dept=department))

# Sửa thông tin nhân sự (Admin sửa tất cả, User chỉ tự sửa chính mình)
@app.route('/admin/edit/<int:id>', methods=['POST'])
def edit_employee(id):
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    current_emp = Employee.query.filter_by(employee_code=session.get('employee_code')).first()
    if not current_emp:
        session.clear()
        return redirect(url_for('login'))

    is_admin = (session.get('role') == 'admin')
    target_emp = Employee.query.get_or_404(id)
    
    # Phân quyền: User thường chỉ được sửa thông tin của CHÍNH MÌNH!
    if not is_admin and target_emp.employee_code != current_emp.employee_code:
        flash('Bạn không có quyền chỉnh sửa thông tin của người khác!', 'danger')
        return redirect(url_for('admin_employees'))
        
    # Cả Admin và User đều được sửa thông tin cá nhân cơ bản
    target_emp.full_name = request.form.get('full_name', target_emp.full_name).strip()
    target_emp.email = request.form.get('email', target_emp.email).strip()
    target_emp.phone = request.form.get('phone', target_emp.phone).strip()
    if 'emergency_phone' in request.form:
        target_emp.emergency_phone = request.form.get('emergency_phone', '').strip()
    
    raw_password = request.form.get('password')
    if raw_password and raw_password.strip():
        target_emp.password_hash = raw_password.strip()
        
    # Riêng Admin mới có quyền đổi các trường hệ thống
    if is_admin:
        new_code = request.form.get('employee_code', target_emp.employee_code).strip()
        if new_code != target_emp.employee_code:
            if Employee.query.filter_by(employee_code=new_code).first():
                flash(f'Mã nhân viên {new_code} đã tồn tại!', 'danger')
                return redirect(url_for('admin_employees', dept=target_emp.department))
            target_emp.employee_code = new_code
            
        target_emp.position = request.form.get('position', target_emp.position).strip()
        target_emp.role = request.form.get('role', target_emp.role)
        target_emp.department = request.form.get('department', target_emp.department)
        start_date_str = request.form.get('start_date')
        if start_date_str:
            try:
                target_emp.start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
            except ValueError:
                pass
                
    db.session.commit()
    
    if target_emp.employee_code == current_emp.employee_code:
        session['full_name'] = target_emp.full_name
        
    flash('Cập nhật thông tin thành công!', 'success')
    tab = 'inactive' if target_emp.status == 'inactive' else 'active'
    return redirect(url_for('admin_employees', dept=target_emp.department, tab=tab))

# Xóa nhân sự (Chuyển vào Đã nghỉ việc / Đã ẩn)
@app.route('/admin/delete/<int:id>')
def delete_employee(id):
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    if session.get('role') != 'admin':
        flash('Chỉ Quản trị viên (Admin) mới có quyền xóa / ẩn nhân sự!', 'danger')
        return redirect(url_for('admin_employees'))
        
    emp = Employee.query.get_or_404(id)
    if emp.employee_code == 'ADMIN01' or emp.employee_code == session.get('employee_code'):
        flash('Không thể xóa tài khoản quản trị viên chính hoặc tài khoản đang đăng nhập!', 'warning')
        return redirect(url_for('admin_employees', dept=emp.department))
        
    dept = emp.department
    emp.status = 'inactive'
    db.session.commit()
    flash(f'Đã chuyển nhân sự {emp.full_name} ({emp.employee_code}) sang danh sách Đã nghỉ việc / Đã ẩn!', 'success')
    return redirect(url_for('admin_employees', dept=dept, tab='inactive'))

# Khôi phục nhân sự về trạng thái Đang làm việc
@app.route('/admin/restore/<int:id>')
def restore_employee(id):
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    if session.get('role') != 'admin':
        flash('Chỉ Quản trị viên (Admin) mới có quyền khôi phục nhân sự!', 'danger')
        return redirect(url_for('admin_employees'))
        
    emp = Employee.query.get_or_404(id)
    dept = emp.department
    emp.status = 'active'
    db.session.commit()
    flash(f'Đã khôi phục thành viên {emp.full_name} ({emp.employee_code}) trở lại danh sách đang làm việc!', 'success')
    return redirect(url_for('admin_employees', dept=dept, tab='active'))

# Xóa vĩnh viễn nhân sự khỏi CSDL (tùy chọn)
@app.route('/admin/permanent-delete/<int:id>')
def permanent_delete_employee(id):
    if not session.get('employee_code'):
        return redirect(url_for('login'))
        
    if session.get('role') != 'admin':
        flash('Chỉ Quản trị viên (Admin) mới có quyền xóa vĩnh viễn!', 'danger')
        return redirect(url_for('admin_employees'))
        
    emp = Employee.query.get_or_404(id)
    if emp.employee_code == 'ADMIN01' or emp.employee_code == session.get('employee_code'):
        flash('Không thể xóa tài khoản quản trị viên chính hoặc tài khoản đang đăng nhập!', 'warning')
        return redirect(url_for('admin_employees', dept=emp.department))
        
    dept = emp.department
    name = emp.full_name
    code = emp.employee_code
    Task.query.filter_by(employee_code=code).delete()
    Report.query.filter_by(employee_code=code).delete()
    LeaveRequest.query.filter_by(employee_code=code).delete()
    db.session.delete(emp)
    db.session.commit()
    flash(f'Đã xóa vĩnh viễn nhân sự {name} ({code}) khỏi cơ sở dữ liệu!', 'success')
    return redirect(url_for('admin_employees', dept=dept, tab='inactive'))

# ==============================================================================
# 7. ROUTE XUẤT EXCEL (TINH GỌN, KHÔNG MÀU MÈ)
# ==============================================================================

@app.route('/admin/export-excel')
def admin_export_excel():
    if not session.get('employee_code') and not app.testing:
        return redirect(url_for('login'))

    dept = request.args.get('dept') or session.get('department') or 'TDH'
    period_type = request.args.get('period_type', 'month')
    month_val = request.args.get('month_val')
    week_start = request.args.get('week_start')

    today = date.today()
    if period_type == 'week' and week_start:
        try:
            start_date = datetime.strptime(week_start, '%Y-%m-%d').date()
        except Exception:
            start_date = today - timedelta(days=6)
        end_date = start_date + timedelta(days=6)
        period_label = f"Tuần {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}"
        file_period_str = f"Tuan_{start_date.strftime('%d%m%Y')}_{end_date.strftime('%d%m%Y')}"
    else:
        period_type = 'month'
        if month_val:
            try:
                parts = month_val.split('-')
                year, month = int(parts[0]), int(parts[1])
            except Exception:
                year, month = today.year, today.month
        else:
            year, month = today.year, today.month
        start_date = date(year, month, 1)
        last_day = calendar.monthrange(year, month)[1]
        end_date = date(year, month, last_day)
        period_label = f"Tháng {month:02d}/{year}"
        file_period_str = f"Thang_{month:02d}_{year}"

    if dept and dept in ['TDH', 'A11A', 'LENS']:
        employees = Employee.query.filter(
            Employee.department == dept,
            (Employee.status != 'inactive') | (Employee.status == None)
        ).all()
        tasks = Task.query.join(Employee, Task.employee_code == Employee.employee_code)\
                          .filter(Employee.department == dept).all()
        reports = Report.query.join(Employee, Report.employee_code == Employee.employee_code)\
                              .filter(Employee.department == dept,
                                      Report.report_date >= start_date,
                                      Report.report_date <= end_date)\
                              .order_by(Report.report_date.desc(), Report.id.desc()).all()
        if not reports:
            reports = Report.query.join(Employee, Report.employee_code == Employee.employee_code)\
                                  .filter(Employee.department == dept)\
                                  .order_by(Report.report_date.desc(), Report.id.desc()).limit(50).all()
    else:
        employees = Employee.query.filter((Employee.status != 'inactive') | (Employee.status == None)).all()
        tasks = Task.query.all()
        reports = Report.query.filter(Report.report_date >= start_date, Report.report_date <= end_date)\
                              .order_by(Report.report_date.desc(), Report.id.desc()).all()
        if not reports:
            reports = Report.query.order_by(Report.report_date.desc(), Report.id.desc()).limit(50).all()

    wb = openpyxl.Workbook()

    # 1. Sheet: Tổng hợp dự án
    ws1 = wb.active
    ws1.title = "Tổng hợp dự án"
    emp_headers = ["STT", "Mã NV", "Họ và Tên", "Phòng Ban", "Chức Vụ", "Tổng CV", "Đang Làm", "Hoàn Thành", "Vướng Mắc", "Tỷ Lệ HT (%)", "Đánh Giá"]
    emp_rows = []
    for idx, e in enumerate(employees, start=1):
        e_tasks = [t for t in tasks if t.employee_code == e.employee_code]
        e_total = len(e_tasks)
        e_prog = sum(1 for t in e_tasks if t.status == 'Đang thực hiện')
        e_done = sum(1 for t in e_tasks if t.status == 'Đã hoàn thành')
        e_stuck = sum(1 for t in e_tasks if t.status in ['Vướng mắc', 'Chưa hoàn thành'])
        e_pct = round(e_done / e_total * 100, 1) if e_total else 0
        ratings = [getattr(t, 'rating', None) for t in e_tasks if getattr(t, 'rating', None) and t.rating != '-']
        e_rating = ratings[0] if ratings else ('Tốt' if e_done > 0 else '—')
        emp_rows.append([
            (idx, EXCEL_ALIGN_CENTER),
            (e.employee_code, EXCEL_ALIGN_CENTER),
            (e.full_name, EXCEL_ALIGN_LEFT),
            (e.department or dept, EXCEL_ALIGN_CENTER),
            (e.position or 'Nhân viên', EXCEL_ALIGN_LEFT),
            (e_total, EXCEL_ALIGN_CENTER),
            (e_prog, EXCEL_ALIGN_CENTER),
            (e_done, EXCEL_ALIGN_CENTER),
            (e_stuck, EXCEL_ALIGN_CENTER),
            (f"{e_pct}%", EXCEL_ALIGN_CENTER),
            (e_rating, EXCEL_ALIGN_CENTER)
        ])
    write_simple_table(
        ws1,
        title=f"BÁO CÁO TỔNG HỢP DỰ ÁN & TIẾN ĐỘ - PHÒNG BAN {dept}",
        subtitle=f"Kỳ báo cáo: {period_label} | Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        headers=emp_headers,
        data_rows=emp_rows
    )

    # 2. Sheet: Chi tiết tiến độ
    ws2 = wb.create_sheet(title="Chi tiết tiến độ")
    task_headers = [
        "STT", "Mã CV", "Tên Công Việc / Dự Án", "Nhân Viên Phụ Trách", 
        "Mã NV", "Phòng Ban", "Hạn Chót", "Mức Độ", "Trạng Thái", 
        "Tiến Độ (%)", "Đánh Giá", "Mô Tả Chi Tiết"
    ]
    task_rows = []
    for idx, t in enumerate(tasks, start=1):
        emp_name = t.assignee.full_name if t.assignee else ''
        dept_name = t.assignee.department if t.assignee else dept
        deadline_str = t.deadline.strftime('%d/%m/%Y') if t.deadline else '—'
        pct_str = f"{t.progress}%" if t.progress is not None else ('100%' if t.status == 'Đã hoàn thành' else '60%')
        task_rows.append([
            (idx, EXCEL_ALIGN_CENTER),
            (t.id, EXCEL_ALIGN_CENTER),
            (t.title or '', EXCEL_ALIGN_LEFT),
            (emp_name, EXCEL_ALIGN_LEFT),
            (t.employee_code, EXCEL_ALIGN_CENTER),
            (dept_name, EXCEL_ALIGN_CENTER),
            (deadline_str, EXCEL_ALIGN_CENTER),
            (t.priority or 'Bình thường', EXCEL_ALIGN_CENTER),
            (t.status or 'Đang thực hiện', EXCEL_ALIGN_CENTER),
            (pct_str, EXCEL_ALIGN_CENTER),
            (t.rating or '—', EXCEL_ALIGN_CENTER),
            (t.description or '—', EXCEL_ALIGN_LEFT)
        ])
    write_simple_table(
        ws2,
        title=f"CHI TIẾT TIẾN ĐỘ CÔNG VIỆC & DỰ ÁN - PHÒNG BAN {dept}",
        subtitle=f"Kỳ báo cáo: {period_label} | Tổng số công việc: {len(tasks)}",
        headers=task_headers,
        data_rows=task_rows
    )

    # 3. Sheet: Báo cáo hàng ngày
    ws3 = wb.create_sheet(title="Báo cáo hàng ngày")
    rep_headers = [
        "STT", "Ngày Báo Cáo", "Thời Gian Gửi", "Mã Nhân Viên", 
        "Họ và Tên Nhân Viên", "Phòng Ban", "Hạng Mục Công Việc", "Nội Dung Báo Cáo Chi Tiết"
    ]
    rep_rows = []
    for idx, r in enumerate(reports, start=1):
        emp = r.employee
        emp_name = emp.full_name if emp else ''
        dept_name = emp.department if emp else dept
        r_date_str = r.report_date.strftime('%d/%m/%Y') if r.report_date else '—'
        r_time_str = r.created_at.strftime('%H:%M:%S') if r.created_at else '—'
        rep_rows.append([
            (idx, EXCEL_ALIGN_CENTER),
            (r_date_str, EXCEL_ALIGN_CENTER),
            (r_time_str, EXCEL_ALIGN_CENTER),
            (r.employee_code, EXCEL_ALIGN_CENTER),
            (emp_name, EXCEL_ALIGN_LEFT),
            (dept_name, EXCEL_ALIGN_CENTER),
            (r.work_types or 'Khác', EXCEL_ALIGN_LEFT),
            (r.description or '—', EXCEL_ALIGN_LEFT)
        ])
    write_simple_table(
        ws3,
        title=f"LỊCH SỬ BÁO CÁO CÔNG VIỆC HÀNG NGÀY - PHÒNG BAN {dept}",
        subtitle=f"Kỳ báo cáo: {period_label} | Tổng số báo cáo ghi nhận: {len(reports)}",
        headers=rep_headers,
        data_rows=rep_rows
    )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    download_filename = f"BaoCao_TongHopDuAn_{dept}_{file_period_str}.xlsx"
    return send_file(
        output,
        download_name=download_filename,
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/admin/export-excel/employee/<string:employee_code>')
@app.route('/export/employee/<string:employee_code>')
def export_employee_excel(employee_code):
    if not session.get('employee_code') and not app.testing:
        return redirect(url_for('login'))

    current_code = session.get('employee_code')
    is_admin = (session.get('role') == 'admin')
    if not is_admin and current_code != employee_code:
        flash('Bạn chỉ có thể xuất báo cáo cá nhân của chính mình!', 'warning')
        return redirect(request.referrer or url_for('admin_dashboard'))
        
    emp = Employee.query.filter_by(employee_code=employee_code).first()
    if not emp:
        flash('Không tìm thấy nhân viên!', 'danger')
        return redirect(url_for('admin_dashboard'))

    wb = openpyxl.Workbook()
    tasks = Task.query.filter_by(employee_code=employee_code).all()
    reports = Report.query.filter_by(employee_code=employee_code).order_by(Report.report_date.desc(), Report.id.desc()).all()

    # 1. Sheet: Công Việc & Dự Án
    ws_tasks = wb.active
    ws_tasks.title = "Công Việc & Dự Án"
    task_headers = ["STT", "Tên Công Việc / Dự Án", "Mức Độ", "Hạn Chót", "Trạng Thái", "Tiến Độ (%)", "Đánh Giá", "Mô Tả Chi Tiết"]
    task_rows = []
    for idx, t in enumerate(tasks, start=1):
        task_rows.append([
            (idx, EXCEL_ALIGN_CENTER),
            (t.title or '', EXCEL_ALIGN_LEFT),
            (t.priority or 'Bình thường', EXCEL_ALIGN_CENTER),
            (t.deadline.strftime('%d/%m/%Y') if t.deadline else '', EXCEL_ALIGN_CENTER),
            (t.status or 'Đang thực hiện', EXCEL_ALIGN_CENTER),
            (f"{t.progress}%" if t.progress is not None else '0%', EXCEL_ALIGN_CENTER),
            (t.rating or '', EXCEL_ALIGN_CENTER),
            (t.description or '', EXCEL_ALIGN_LEFT)
        ])
    write_simple_table(
        ws_tasks,
        title=f"BÁO CÁO CÔNG VIỆC & DỰ ÁN - {emp.full_name.upper()} ({emp.employee_code})",
        subtitle=f"Phòng ban: {emp.department or 'TDH'} | Chức vụ: {emp.position or 'Nhân viên'} | Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        headers=task_headers,
        data_rows=task_rows
    )

    # 2. Sheet: Báo Cáo Hàng Ngày
    ws_reports = wb.create_sheet(title="Báo Cáo Hàng Ngày")
    report_headers = ["STT", "Ngày Báo Cáo", "Hạng Mục", "Nội Dung Báo Cáo Chi Tiết", "Thời Gian Gửi"]
    rep_rows = []
    for idx, rep in enumerate(reports, start=1):
        rep_rows.append([
            (idx, EXCEL_ALIGN_CENTER),
            (rep.report_date.strftime('%d/%m/%Y') if rep.report_date else '', EXCEL_ALIGN_CENTER),
            (rep.work_types or '', EXCEL_ALIGN_LEFT),
            (rep.description or '', EXCEL_ALIGN_LEFT),
            (rep.created_at.strftime('%d/%m/%Y %H:%M') if rep.created_at else '', EXCEL_ALIGN_CENTER)
        ])
    write_simple_table(
        ws_reports,
        title=f"LỊCH SỬ BÁO CÁO HÀNG NGÀY - {emp.full_name.upper()} ({emp.employee_code})",
        subtitle=f"Phòng ban: {emp.department or 'TDH'} | Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        headers=report_headers,
        data_rows=rep_rows
    )

    # 3. Sheet: Thông Tin Cá Nhân
    ws_info = wb.create_sheet(title="Thông Tin Cá Nhân")
    ws_info.cell(row=1, column=1, value=f"HỒ SƠ NHÂN VIÊN - {emp.full_name.upper()}").font = EXCEL_FONT_TITLE
    info_fields = [
        ("Mã Nhân Viên", emp.employee_code),
        ("Họ và Tên", emp.full_name),
        ("Phòng Ban", emp.department or 'TDH'),
        ("Chức Vụ", emp.position or 'Nhân viên'),
        ("Vai Trò Hệ Thống", "Quản trị viên" if emp.role == 'admin' else 'Nhân viên'),
        ("Email", emp.email or 'Chưa cập nhật'),
        ("Số Điện Thoại", emp.phone or 'Chưa cập nhật'),
        ("Số Khẩn Cấp", emp.emergency_phone or 'Chưa cập nhật'),
        ("Ngày Vào Làm", emp.start_date.strftime('%d/%m/%Y') if emp.start_date else 'Chưa cập nhật'),
        ("Trạng Thái Hoạt Động", "Đang làm việc" if emp.status == 'active' else (emp.status or 'active')),
        ("Tổng Số Công Việc", len(tasks)),
        ("Tổng Số Báo Cáo", len(reports)),
    ]
    for f_idx, (f_name, f_val) in enumerate(info_fields, start=3):
        c1 = ws_info.cell(row=f_idx, column=1, value=f_name)
        c1.font = EXCEL_FONT_HEADER
        c1.border = EXCEL_BORDER
        c1.alignment = EXCEL_ALIGN_LEFT
        c2 = ws_info.cell(row=f_idx, column=2, value=str(f_val))
        c2.font = EXCEL_FONT_DATA
        c2.border = EXCEL_BORDER
        c2.alignment = EXCEL_ALIGN_LEFT
    ws_info.column_dimensions['A'].width = 24
    ws_info.column_dimensions['B'].width = 40

    # Save to buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Naming format: BaoCaoCaNhan_tenthanhvien_ngay-thang-nam.xlsx
    clean_name = re.sub(r'[^a-zA-Z0-9]', '', remove_accents(emp.full_name))
    if not clean_name:
        clean_name = re.sub(r'[^a-zA-Z0-9]', '', remove_accents(emp.employee_code)) or 'ThanhVien'
    today_str = datetime.now().strftime('%d-%m-%Y')
    prefix = request.args.get('prefix', 'BaoCaoCaNhan')
    download_filename = f"{prefix}_{clean_name}_{today_str}.xlsx"
    
    return send_file(
        output,
        download_name=download_filename,
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# ==============================================================================
# 7.1. API LẤY BÁO CÁO HOẠT ĐỘNG CỦA NHÂN VIÊN CHO ADMIN (MODAL XEM BÁO CÁO)
# ==============================================================================
@app.route('/api/admin/reports')
def api_admin_reports():
    if not session.get('employee_code'):
        return {'success': False, 'message': 'Chưa đăng nhập'}, 401
        
    current_emp = Employee.query.filter_by(employee_code=session.get('employee_code')).first()
    if not current_emp:
        return {'success': False, 'message': 'Không tìm thấy người dùng'}, 404

    is_admin = (session.get('role') == 'admin')
    user_dept = current_emp.department or 'TDH'
    
    # Bộ lọc phòng ban (chỉ 3 phòng TDH, A11A, LENS)
    if is_admin:
        dept = request.args.get('dept', 'TDH')
        if dept not in ['TDH', 'A11A', 'LENS']:
            dept = 'TDH'
    else:
        dept = user_dept

    date_str = request.args.get('date')
    view_mode = request.args.get('view', 'day')
    employee_code = request.args.get('employee_code', '').strip()

    # Xử lý ngày
    today = datetime.now().date()
    target_date = today
    if date_str:
        try:
            target_date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            target_date = today

    # Tính toán khoảng ngày theo ngày / tuần
    if view_mode == 'week':
        # Thứ 2 đến Chủ nhật của tuần chứa target_date
        start_date = target_date - timedelta(days=target_date.weekday())
        end_date = start_date + timedelta(days=6)
        query = Report.query.filter(Report.report_date >= start_date, Report.report_date <= end_date)
    else:
        start_date = target_date
        end_date = target_date
        query = Report.query.filter(Report.report_date == target_date)

    # Lọc theo phòng ban (Chỉ xem báo cáo của user phòng đó, tuyệt đối không xem chung)
    query = query.join(Employee, Report.employee_code == Employee.employee_code)\
                 .filter(Employee.department == dept)

    # Lọc theo nhân viên cụ thể (nếu có)
    if employee_code:
        query = query.filter(Report.employee_code == employee_code)

    reports = query.order_by(Report.report_date.desc(), Report.created_at.desc(), Report.id.desc()).all()

    report_list = []
    for r in reports:
        emp = r.employee
        report_list.append({
            'id': r.id,
            'employee_code': r.employee_code,
            'full_name': emp.full_name if emp else r.employee_code,
            'position': emp.position if emp else 'staff',
            'department': emp.department if emp else '',
            'avatar_color': get_avatar_color(r.employee_code + (emp.full_name if emp else '')),
            'initials': get_initials(emp.full_name if emp else r.employee_code),
            'report_date': r.report_date.strftime('%Y-%m-%d'),
            'report_date_formatted': r.report_date.strftime('%d/%m/%Y'),
            'created_time': r.created_at.strftime('%H:%M - %d/%m/%Y') if r.created_at else '',
            'time_only': r.created_at.strftime('%H:%M') if r.created_at else '00:00',
            'work_types': r.work_types or 'Khác',
            'description': r.description or ''
        })

    return {
        'success': True,
        'reports': report_list,
        'dept': dept,
        'date': target_date.strftime('%Y-%m-%d'),
        'view': view_mode,
        'week_start': start_date.strftime('%d/%m/%Y'),
        'week_end': end_date.strftime('%d/%m/%Y'),
        'count': len(report_list)
    }

# ==============================================================================
# 7.2. ROUTE ĐIỀU HƯỚNG BÁO CÁO VỀ DASHBOARD
# ==============================================================================
@app.route('/admin/reports')
def admin_reports():
    dept = request.args.get('dept', 'TDH')
    return redirect(url_for('admin_dashboard', dept=dept))

# ==============================================================================
# 8. ROUTE ĐĂNG KÝ LỊCH NGHỈ PHÉP (DÙNG CHUNG ADMIN & USER)
# ==============================================================================

@app.route('/leave-calendar')
def leave_calendar():
    emp_code = session.get('employee_code')
    if not emp_code:
        flash('Vui lòng đăng nhập trước khi tiếp tục.', 'warning')
        return redirect(url_for('login'))
        
    current_user = Employee.query.filter_by(employee_code=emp_code).first()
    if not current_user:
        session.clear()
        return redirect(url_for('login'))
        
    today = datetime.now().date()
    month = request.args.get('month', type=int) or today.month
    year = request.args.get('year', type=int) or today.year
    is_admin = (session.get('role') == 'admin')
    
    if is_admin:
        selected_dept = request.args.get('dept', 'TDH')
        if selected_dept not in ['TDH', 'A11A', 'LENS']:
            selected_dept = 'TDH'
    else:
        # Nhân viên thường: Cố định theo đúng phòng ban của mình
        user_dept = current_user.department or 'TDH'
        requested_dept = request.args.get('dept')
        if requested_dept and requested_dept != user_dept:
            flash(f'Bạn chỉ có quyền xem lịch nghỉ phép của phòng ban mình ({user_dept})!', 'info')
            return redirect(url_for('leave_calendar', month=month, year=year, dept=user_dept))
        selected_dept = user_dept

    # Tính toán tháng trước / tháng sau
    prev_month, prev_year = (12, year - 1) if month == 1 else (month - 1, year)
    next_month, next_year = (1, year + 1) if month == 12 else (month + 1, year)

    # Lấy toàn bộ đơn nghỉ phép trong tháng này từ CSDL
    all_leaves = LeaveRequest.query.all()
    month_leaves = [
        l for l in all_leaves 
        if l.leave_date and l.leave_date.month == month and l.leave_date.year == year
    ]

    # Lọc đơn nghỉ phép theo đúng phòng ban (TDH, A11A, LENS)
    display_leaves = [l for l in month_leaves if l.employee and l.employee.department == selected_dept]

    # Nhóm danh sách nghỉ theo từng ngày dạng YYYY-MM-DD
    leaves_by_date = {}
    for l in display_leaves:
        date_key = l.leave_date.strftime('%Y-%m-%d')
        if date_key not in leaves_by_date:
            leaves_by_date[date_key] = []
        leaves_by_date[date_key].append(l)

    # Xây dựng ma trận lưới lịch tháng
    cal = calendar.Calendar(firstweekday=0)
    month_dates = cal.monthdatescalendar(year, month)

    calendar_weeks = []
    for week in month_dates:
        week_data = []
        for d in week:
            date_str = d.strftime('%Y-%m-%d')
            week_data.append({
                'day': d.day,
                'date_str': date_str,
                'is_current_month': (d.month == month),
                'is_today': (d == today),
                'is_sunday': (d.isoweekday() == 7),
                'leaves': leaves_by_date.get(date_str, [])
            })
        calendar_weeks.append(week_data)

    # Thống kê số ngày nghỉ của từng nhân viên cho Modal
    emp_stats = {}
    for l in display_leaves:
        emp_name = l.employee.full_name if l.employee else l.employee_code
        dept = l.employee.department if l.employee else ''
        if emp_name not in emp_stats:
            emp_stats[emp_name] = {'count': 0, 'dept': dept, 'dates': []}
        emp_stats[emp_name]['count'] += 1
        emp_stats[emp_name]['dates'].append(l.leave_date.strftime('%d/%m'))

    chart_labels = list(emp_stats.keys())
    chart_data = [info['count'] for info in emp_stats.values()]
    emp_summary = [
        {'name': k, 'dept': v['dept'], 'count': v['count'], 'dates_str': ', '.join(v['dates'])} 
        for k, v in emp_stats.items()
    ]

    user_initials = get_initials(current_user.full_name)

    return render_template(
        'leave_calendar.html',
        current_user=current_user,
        user_initials=user_initials,
        calendar_weeks=calendar_weeks,
        current_month=month,
        current_year=year,
        prev_month=prev_month,
        prev_year=prev_year,
        next_month=next_month,
        next_year=next_year,
        selected_dept=selected_dept,
        chart_labels=chart_labels,
        chart_data=chart_data,
        emp_summary=emp_summary,
        is_admin=is_admin,
        active_page='leave_calendar'
    )

# Đăng ký nghỉ phép
@app.route('/leave/register', methods=['POST'])
def leave_register():
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))
        
    dates_str = request.form.get('leave_dates', '')
    reason = request.form.get('reason', '').strip()
    redirect_month = request.form.get('redirect_month', type=int)
    redirect_year = request.form.get('redirect_year', type=int)
    redirect_dept = request.form.get('redirect_dept', '')
    
    emp_user = Employee.query.filter_by(employee_code=emp_code).first()
    is_admin = (session.get('role') == 'admin')
    if not is_admin and emp_user:
        redirect_dept = emp_user.department or 'TDH'

    if not dates_str:
        flash('Vui lòng chọn ít nhất một ngày trên lịch để đăng ký!', 'warning')
        return redirect(url_for('leave_calendar', month=redirect_month, year=redirect_year, dept=redirect_dept))

    if not reason:
        flash('Vui lòng nhập lý do nghỉ phép!', 'warning')
        return redirect(url_for('leave_calendar', month=redirect_month, year=redirect_year, dept=redirect_dept))

    date_list = [d.strip() for d in dates_str.split(',') if d.strip()]
    added_count = 0

    for d_str in date_list:
        try:
            leave_date = datetime.strptime(d_str, '%Y-%m-%d').date()
            exist = LeaveRequest.query.filter_by(employee_code=emp_code, leave_date=leave_date).first()
            if not exist:
                new_leave = LeaveRequest(
                    employee_code=emp_code,
                    leave_date=leave_date,
                    reason=reason
                )
                db.session.add(new_leave)
                added_count += 1
        except ValueError:
            continue

    db.session.commit()
    if added_count > 0:
        flash(f'Đã đăng ký thành công {added_count} ngày nghỉ phép!', 'success')
    else:
        flash('Các ngày bạn chọn đã được đăng ký trước đó!', 'info')

    return redirect(url_for('leave_calendar', month=redirect_month, year=redirect_year, dept=redirect_dept))

# Hủy lịch nghỉ phép
@app.route('/leave/cancel/<int:id>', methods=['POST'])
def leave_cancel(id):
    emp_code = session.get('employee_code')
    if not emp_code:
        return redirect(url_for('login'))

    leave = LeaveRequest.query.get_or_404(id)
    if leave.employee_code == emp_code or session.get('role') == 'admin':
        month = leave.leave_date.month
        year = leave.leave_date.year
        db.session.delete(leave)
        db.session.commit()
        flash('Đã hủy lịch nghỉ phép thành công!', 'info')
        return redirect(url_for('leave_calendar', month=month, year=year))

    flash('Bạn không có quyền hủy lịch nghỉ này!', 'danger')
    return redirect(url_for('leave_calendar'))

# ==============================================================================
# 9. KHỞI CHẠY ỨNG DỤNG
# ==============================================================================
if __name__ == '__main__':
    with app.app_context():
        db.create_all()

        # Đảm bảo luôn có ít nhất 1 tài khoản Admin mặc định nếu database trống
        admin_exists = Employee.query.filter_by(role='admin').first()
        if not admin_exists:
            admin_user = Employee(
                employee_code='ADMIN01',
                password_hash='123456',
                role='admin',
                full_name='Quản Trị Viên',
                email='admin@example.com',
                phone='0900000000',
                position='Quản lý hệ thống',
                department='TDH'
            )
            db.session.add(admin_user)
            db.session.commit()
            print("Đã tạo tài khoản Admin mặc định (ADMIN01 / 123456) thành công!")

        # Tạo sẵn dữ liệu nhóm dự án chung mẫu (theo ảnh mẫu) nếu chưa có
        if ProjectGroup.query.count() == 0:
            sample_groups = [
                {'name': 'Máy kiểm tra sản phẩm cho cty Denso', 'creator': 'Vũ Hữu Việt', 'members': 'Vũ Hữu Việt', 'count': 1},
                {'name': 'Dự án hút nhả chân pin đồng chuyền Nex', 'creator': 'Bùi Văn Nhân', 'members': 'Nguyễn Viết Diệp, Lê Anh Dũng, Bùi Văn Nhân', 'count': 3},
                {'name': 'Dự án jig kẹp 32 sản phẩm y tế', 'creator': 'Bùi Văn Nhân', 'members': 'Nguyễn Viết Diệp, Bùi Văn Nhân', 'count': 3},
                {'name': 'Dự án máy check bavia F3', 'creator': 'Bùi Văn Nhân', 'members': 'Nguyễn Viết Diệp, Lê Anh Dũng, Bùi Văn Nhân', 'count': 3},
                {'name': 'Phần mềm Checksheet F2', 'creator': 'Phạm Quang Trọng', 'members': 'Phạm Quang Trọng', 'count': 5},
            ]
            for g in sample_groups:
                pg = ProjectGroup(
                    name=g['name'],
                    department='TDH',
                    creator_name=g['creator'],
                    members_str=g['members'],
                    task_count=g['count']
                )
                db.session.add(pg)
            db.session.commit()
            print("Đã khởi tạo các nhóm dự án chung mẫu thành công!")

    app.run(debug=True, port=5000)
