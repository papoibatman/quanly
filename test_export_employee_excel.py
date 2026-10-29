import io
import os
import sys
import re
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')
sys.stderr.reconfigure(encoding='utf-8')

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from app import app, db, Employee, Task, Report

def test_export():
    app.config['TESTING'] = True
    client = app.test_client()

    with app.app_context():
        # Find an employee
        emp = Employee.query.first()
        assert emp is not None, "No employee found in db"
        print(f"Testing employee: {emp.full_name} ({emp.employee_code})")

        # Test unaccented name logic
        from app import remove_accents
        unaccented = remove_accents(emp.full_name)
        clean_name = re.sub(r'[^a-zA-Z0-9]', '', unaccented)
        today_str = datetime.now().strftime('%d-%m-%Y')
        expected_filename = f"BaoCaoCaNhan_{clean_name}_{today_str}.xlsx"
        print(f"Expected filename: {expected_filename}")

        # Simulate logged in user
        with client.session_transaction() as sess:
            sess['employee_code'] = 'ADMIN01'
            sess['role'] = 'admin'

        # Test endpoint
        resp = client.get(f"/admin/export-excel/employee/{emp.employee_code}")
        print(f"Response status: {resp.status_code}")
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"

        # Check Content-Disposition header
        cd = resp.headers.get('Content-Disposition', '')
        print(f"Content-Disposition: {cd}")
        assert expected_filename in cd, f"Expected {expected_filename} in Content-Disposition"
        assert "BaoCaoCaNhan" in cd, "Expected 'BaoCaoCaNhan' in Content-Disposition"

        # Check content-type
        ct = resp.headers.get('Content-Type', '')
        assert 'spreadsheetml' in ct or 'application/vnd.openxmlformats' in ct, f"Wrong Content-Type: {ct}"

        # Parse with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        sheet_names = wb.sheetnames
        print(f"Workbook sheets: {sheet_names}")
        assert "Công Việc & Dự Án" in sheet_names, "Missing sheet 'Công Việc & Dự Án'"
        assert "Báo Cáo Hàng Ngày" in sheet_names, "Missing sheet 'Báo Cáo Hàng Ngày'"
        assert "Thông Tin Cá Nhân" in sheet_names, "Missing sheet 'Thông Tin Cá Nhân'"

        # Verify ws_tasks
        ws_tasks = wb["Công Việc & Dự Án"]
        assert ws_tasks['A1'].value is not None
        print(f"Sheet 1 title cell: {ws_tasks['A1'].value}")
        print(f"Sheet 1 headers (row 4): {[ws_tasks.cell(row=4, column=i).value for i in range(1, 9)]}")

        # Verify ws_reports
        ws_reports = wb["Báo Cáo Hàng Ngày"]
        print(f"Sheet 2 title cell: {ws_reports['A1'].value}")
        print(f"Sheet 2 headers (row 4): {[ws_reports.cell(row=4, column=i).value for i in range(1, 6)]}")

        # Verify ws_info
        ws_info = wb["Thông Tin Cá Nhân"]
        print(f"Sheet 3 title cell: {ws_info['A1'].value}")
        print(f"Sheet 3 info field 1: {ws_info['A3'].value} = {ws_info['B3'].value}")

        # Test second route alias
        resp2 = client.get(f"/export/employee/{emp.employee_code}")
        assert resp2.status_code == 200
        assert expected_filename in resp2.headers.get('Content-Disposition', '')
        print("Second route alias /export/employee/<code_pqt> also verified OK!")

        # Test multiple employees with Vietnamese accents
        emps = Employee.query.limit(5).all()
        for e in emps:
            e_unaccented = remove_accents(e.full_name)
            e_clean = re.sub(r'[^a-zA-Z0-9]', '', e_unaccented) or 'NV'
            e_expected = f"BaoCaoCaNhan_{e_clean}_{today_str}.xlsx"
            e_resp = client.get(f"/admin/export-excel/employee/{e.employee_code}")
            assert e_resp.status_code == 200
            assert e_expected in e_resp.headers.get('Content-Disposition', '')
            print(f"Verified {e.full_name} -> {e_expected}")

        print("\nALL EXPORT EXCEL TESTS PASSED SUCCESSFULLY!")

if __name__ == '__main__':
    test_export()
